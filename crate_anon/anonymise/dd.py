#!/usr/bin/env python

"""
crate_anon/anonymise/dd.py

===============================================================================

    Copyright (C) 2015-2021 Rudolf Cardinal (rudolf@pobox.com).

    This file is part of CRATE.

    CRATE is free software: you can redistribute it and/or modify
    it under the terms of the GNU General Public License as published by
    the Free Software Foundation, either version 3 of the License, or
    (at your option) any later version.

    CRATE is distributed in the hope that it will be useful,
    but WITHOUT ANY WARRANTY; without even the implied warranty of
    MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE. See the
    GNU General Public License for more details.

    You should have received a copy of the GNU General Public License
    along with CRATE. If not, see <https://www.gnu.org/licenses/>.

===============================================================================

**Data dictionary classes for CRATE anonymiser.**

The data dictionary is a TSV file, for ease of editing by multiple authors,
rather than a database table.

"""

# =============================================================================
# Imports
# =============================================================================

from collections import Counter, OrderedDict
import csv
from dataclasses import dataclass
from functools import lru_cache
from itertools import zip_longest
import logging
import operator
import os
from typing import (
    AbstractSet, Any, Callable, Dict, Iterable, List, Optional,
    TextIO, Tuple, TYPE_CHECKING, Union
)

from cardinal_pythonlib.file_io import smart_open
from cardinal_pythonlib.sql.validation import is_sqltype_integer
from cardinal_pythonlib.sqlalchemy.schema import (
    is_sqlatype_integer,
    is_sqlatype_string,
    is_sqlatype_text_over_one_char,
)
import openpyxl
import pyexcel_ods
import pyexcel_xlsx
from sortedcontainers import SortedSet
import sqlalchemy.exc
from sqlalchemy import Column, Table, DateTime
from sqlalchemy.engine.interfaces import Dialect
from sqlalchemy.sql.sqltypes import String, TypeEngine

# don't import config: circular dependency would have to be sorted out
from crate_anon.anonymise.constants import (
    AlterMethodType,
    AnonymiseConfigKeys,
    TABLE_KWARGS,
    ScrubMethod,
    ScrubSrc,
    SrcFlag,
    TridType,
)
from crate_anon.anonymise.ddr import DataDictionaryRow
from crate_anon.anonymise.scrub import PersonalizedScrubber

if TYPE_CHECKING:
    from crate_anon.anonymise.config import Config

log = logging.getLogger(__name__)


# =============================================================================
# Constants
# =============================================================================

EXT_TSV = ".tsv"
EXT_ODS = ".ods"
EXT_XLSX = ".xlsx"

STRING_LENGTH_FOR_BIGINT = len(str(-2 ** 63))
# = -2^63: https://dev.mysql.com/doc/refman/8.0/en/integer-types.html


# =============================================================================
# Helper classes
# =============================================================================

@dataclass
class ScrubSourceFieldInfo:
    is_mpid: bool
    is_patient: bool
    recurse: bool
    required_scrubber: bool
    scrub_method: ScrubMethod
    signature: str
    value_fieldname: str


# =============================================================================
# Helper functions
# =============================================================================

def ensure_no_source_type_mismatch(ddr: DataDictionaryRow,
                                   config_sqlatype: Union[TypeEngine, String],
                                   primary_pid: bool = True) -> None:
    """
    Ensure that the source column type of a data dictionary row is compatible
    with what's expected from the config. We check this only for specific type
    of column (PID, MPID), because we need to know their data types concretely
    for the secret mapping table. The question is not whether the types are the
    same, but whether the value will fit into the config-determined type (for
    example, it's OK to convert an integer to a long-enough string but
    necessarily not the other way round).

    Args:
        ddr:
            Data dictionary row.
        config_sqlatype:
            SQLAlchemy column type that would be expected based on the current
            config.
        primary_pid:
            Is this the main PID field? If false, it's the MPID.
    """
    if primary_pid:
        human_type = "primary PID"
        configparam = AnonymiseConfigKeys.SQLATYPE_PID
    else:
        human_type = "master PID"
        configparam = AnonymiseConfigKeys.SQLATYPE_MPID
    rowtype = ddr.src_sqla_coltype
    suffix = ""
    if is_sqlatype_integer(rowtype):
        # ---------------------------------------------------------------------
        # Integer source
        # ---------------------------------------------------------------------
        if is_sqlatype_integer(config_sqlatype):
            # Good enough. The only integer type we use for PID/MPID is
            # BigInteger, so any integer type should fit.
            return
        elif is_sqlatype_string(config_sqlatype):
            # Storing an integer in a string. This may be OK, if the string is
            # long enough. We could do detailed checks here based on the type
            # of integer, but we'll be simple.
            if STRING_LENGTH_FOR_BIGINT <= config_sqlatype.length:
                # It'll fit!
                return
            else:
                suffix = (
                    f"Using a bigger string field in the config (minimum "
                    f"length {STRING_LENGTH_FOR_BIGINT}) would fix this."
                )
    elif is_sqlatype_string(rowtype):
        # ---------------------------------------------------------------------
        # String source
        # ---------------------------------------------------------------------
        # Strings are fine if we will store them in a long-enough string.
        if is_sqlatype_string(config_sqlatype):
            # noinspection PyUnresolvedReferences
            if rowtype.length <= config_sqlatype.length:
                return
            else:
                suffix = (
                    f"Using a bigger string field in the config (minimum "
                    f"length {rowtype.length}) would fix this."
                )
    # Generic error:
    raise ValueError(
        f"Source column {ddr.src_signature} is marked as a "
        f"{human_type} field but its type is {rowtype}, "
        f"while the config thinks it should be {config_sqlatype} "
        f"(determined by the {configparam!r} parameter). {suffix}")


# =============================================================================
# DataDictionary
# =============================================================================

class DataDictionary(object):
    """
    Class representing an entire data dictionary.
    """

    def __init__(self, config: "Config") -> None:
        """
        Set defaults.

        Args:
            config: :class:`crate_anon.anonymise.config.Config`
        """
        self.config = config
        self.rows = []  # type: List[DataDictionaryRow]
        # noinspection PyArgumentList
        self.cached_srcdb_table_pairs = SortedSet()

    # -------------------------------------------------------------------------
    # Information
    # -------------------------------------------------------------------------

    @property
    def n_rows(self) -> int:
        """
        Number of rows.
        """
        return len(self.rows)

    # -------------------------------------------------------------------------
    # Loading
    # -------------------------------------------------------------------------

    def read_from_file(self,
                       filename: str,
                       check_valid: bool = True,
                       override_dialect: Dialect = None) -> None:
        """
        Read DD from file.

        Args:
            filename:
                Filename to read.
            check_valid:
                Run a validity check after setting each row from its values?
            override_dialect:
                SQLAlchemy SQL dialect to enforce (e.g. for interpreting
                textual column types in the source database). By default, the
                source database's own dialect is used.
        """
        log.debug(f"Loading data dictionary: {filename}")
        _, ext = os.path.splitext(filename)
        if ext == EXT_TSV:
            row_gen = self._gen_rows_from_tsv(filename)
        elif ext == EXT_ODS:
            row_gen = self._gen_rows_from_ods(filename)
        elif ext == EXT_XLSX:
            row_gen = self._gen_rows_from_xlsx(filename)
        else:
            raise ValueError(f"Unknown data dictionary extension: {ext!r}")
        self._read_from_rows(row_gen,
                             check_valid=check_valid,
                             override_dialect=override_dialect)

    @staticmethod
    def _skip_row(row: List[Any]) -> bool:
        """
        Should we skip a row, because it's empty or starts with a comment?
        """
        if not row:
            return True
        first = row[0]
        if isinstance(first, str) and first.strip().startswith("#"):
            return True
        return not any(v for v in row)

    @classmethod
    def _gen_rows_from_tsv(cls, filename: str) -> Iterable[List[Any]]:
        """
        Generates rows from a TSV file.
        """
        log.debug(f"Loading as TSV: {filename}")
        with open(filename, 'r') as tsvfile:
            tsv = csv.reader(tsvfile, delimiter='\t')
            for row in tsv:
                if cls._skip_row(row):
                    continue
                yield row

    @classmethod
    def _gen_rows_from_xlsx(cls, filename: str) -> Iterable[List[Any]]:
        """
        Generates rows from an XLSX file, reading the first sheet.
        """
        log.debug(f"Loading as XLSX: {filename}")
        workbook = openpyxl.load_workbook(filename)
        # ... NB potential bug using read_only; see postcodes.py
        worksheet = workbook.active  # first sheet, by default
        for sheet_row in worksheet.iter_rows():
            row = [
                "" if cell.value is None else cell.value
                for cell in sheet_row
            ]
            if cls._skip_row(row):
                continue
            yield row

    @classmethod
    def _gen_rows_from_ods(cls, filename: str) -> Iterable[List[Any]]:
        """
        Generates rows from an ODS file, reading the first sheet.
        """
        log.debug(f"Loading as ODS: {filename}")
        data = pyexcel_ods.get_data(filename)  # type: Dict[str, List[List[Any]]]  # noqa
        # ... but it's an ordered dictionary, so:
        first_key = next(iter(data))
        first_sheet_rows = data[first_key]
        for row in first_sheet_rows:
            if cls._skip_row(row):
                continue
            yield row

    def _read_from_rows(self,
                        rows: Iterable[List[Any]],
                        check_valid: bool = True,
                        override_dialect: Dialect = None) -> None:
        """
        Internal function to read from a set of rows, whatever the underlying
        format.

        Args:
            rows:
                Iterable of rows (one per data dictionary row), each row being
                a list of values.
            check_valid:
                Run a validity check after setting the values?
            override_dialect:
                SQLAlchemy SQL dialect to enforce (e.g. for interpreting
                textual column types in the source database). By default, the
                source database's own dialect is used.
        """
        # Clear existing data
        self.rows = []  # type: List[DataDictionaryRow]

        # Headers
        # noinspection PyTypeChecker
        headers = next(rows)
        if not all(x in headers for x in DataDictionaryRow.ROWNAMES):
            actual = "\n".join(
                f"{i}. {h}"
                for i, h in enumerate(headers, start=1)
            )
            desired = "\n".join(
                f"{i}. {h}"
                for i, h in enumerate(DataDictionaryRow.ROWNAMES, start=1)
            )
            raise ValueError(
                f"Bad data dictionary file. Data dictionaries must be in "
                f"tabular format and contain the following headings:\n\n"
                f"{desired}\n\n"
                f"but yours are:\n\n"
                f"{actual}"
            )
        log.debug("Data dictionary has correct header. Loading content...")

        # Data
        for values in rows:
            if len(values) < len(headers):
                valuedict = dict(zip_longest(headers, values, fillvalue=""))
            else:
                valuedict = dict(zip(headers, values))
            ddr = DataDictionaryRow(self.config)
            try:
                ddr.set_from_dict(valuedict, override_dialect=override_dialect)
                if check_valid:
                    ddr.check_valid()
            except ValueError:
                log.critical(f"Offending input: {valuedict}")
                raise
            self.rows.append(ddr)
        log.debug("... content loaded.")

        # Clear caches
        self.clear_caches()

    @classmethod
    def create_from_file(cls,
                         filename: str,
                         config: "Config",
                         check_valid: bool = True,
                         override_dialect: Dialect = None) -> "DataDictionary":
        """
        Creates a new data dictionary by reading a file.
        """
        dd = DataDictionary(config)
        dd.read_from_file(filename,
                          check_valid=check_valid,
                          override_dialect=override_dialect)
        return dd

    def draft_from_source_databases(self, report_every: int = 100) -> None:
        """
        Create a draft DD from a source database.

        Will skip any rows it knows about already (thus allowing the generation
        of incremental changes).

        Args:
            report_every: report to the Python log every *n* columns
        """
        log.info("Reading information for draft data dictionary")

        # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
        # Scan databases
        # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
        existing_signatures = set(ddr.src_signature for ddr in self.rows)
        for pretty_dbname, db in self.config.sources.items():
            log.info(f"... database nice name = {pretty_dbname}")
            cfg = db.srccfg
            meta = db.metadata
            i = 0
            # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
            # Scan each table
            # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
            for t in meta.sorted_tables:
                tablename = t.name
                log.info(f"... ... table: {tablename}")
                new_rows = []  # type: List[DataDictionaryRow]
                is_patient_table = False

                # Skip table?
                if cfg.is_table_denied(tablename):
                    log.debug(f"Skipping denied table: {tablename}")
                    continue
                all_col_names = [c.name for c in t.columns]
                if cfg.does_table_fail_minimum_fields(all_col_names):
                    log.debug(f"Skipping table {t} because it fails "
                              f"minimum field requirements")
                    continue

                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                # Scan each column
                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                for c in t.columns:
                    i += 1
                    if report_every and i % report_every == 0:
                        log.debug(f"... reading source field number {i}")
                    columnname = c.name
                    # import pdb; pdb.set_trace()
                    # log.critical(f"str(coltype) == {str(c.type)}")
                    # log.critical(f"repr(coltype) == {repr(c.type)}")
                    try:
                        datatype_sqltext = str(c.type)
                    except sqlalchemy.exc.CompileError:
                        log.critical(f"Column that failed was: {c!r}")
                        raise
                    sqla_coltype = c.type
                    # Do not manipulate the case of SOURCE tables/columns.
                    # If you do, they can fail to match the SQLAlchemy
                    # introspection and cause a crash.
                    # Changed to be a destination manipulation (2016-06-04).
                    if cfg.is_field_denied(columnname):
                        log.debug(f"Skipping denied column: "
                                  f"{tablename}.{columnname}")
                        continue
                    comment = ''  # currently unsupported by SQLAlchemy
                    if cfg.ddgen_append_source_info_to_comment:
                        comment = f"[from {tablename}.{columnname}]"
                    ddr = DataDictionaryRow(self.config)
                    ddr.set_from_src_db_info(
                        pretty_dbname, tablename, columnname,
                        datatype_sqltext,
                        sqla_coltype,
                        dbconf=cfg,
                        comment=comment)

                    # ---------------------------------------------------------
                    # If we have this one already, skip ASAP
                    # This is how incremental data dictionaries get generated.
                    # ---------------------------------------------------------
                    sig = ddr.src_signature
                    if sig in existing_signatures:
                        log.debug(f"Skipping duplicated column: "
                                  f"{tablename}.{columnname}")
                        continue
                    existing_signatures.add(sig)

                    if ddr.contains_patient_info:
                        is_patient_table = True

                    # Checking validity slows us down, and we are after all
                    # creating these programmatically!
                    # ddr.check_valid(self.config)

                    new_rows.append(ddr)

                # Now, table-wide checks across all columns:
                if not is_patient_table:
                    for ddr in new_rows:
                        ddr.remove_scrub_from_alter_methods()
                        # Pointless to scrub in a non-patient table

                self.rows.extend(new_rows)

        log.info("... done")
        self.clear_caches()
        log.info("Revising draft data dictionary")
        for ddr in self.rows:
            if ddr.from_file:
                continue
            # Don't scrub_in non-patient tables
            if (ddr.src_table
                    not in self.get_src_tables_with_patient_info(ddr.src_db)):
                ddr._scrub = False
        log.info("... done")
        self.sort()

    def make_dest_datatypes_explicit(self) -> None:
        """
        By default, when autocreating a data dictionary, the ``dest_datatype``
        field is not populated explicit, just implicitly. This option makes
        them explicit by instantiating those values. Primarily for debugging.
        """
        for ddr in self.rows:
            ddr.make_dest_datatype_explicit()

    # -------------------------------------------------------------------------
    # Sorting
    # -------------------------------------------------------------------------

    def sort(self) -> None:
        """
        Sorts the data dictionary.
        """
        log.info("Sorting data dictionary")
        self.rows = sorted(
            self.rows,
            key=operator.attrgetter("src_db_lowercase",
                                    "src_table_lowercase",
                                    "src_field_lowercase"))
        log.info("... done")

    # -------------------------------------------------------------------------
    # Validation
    # -------------------------------------------------------------------------

    def check_against_source_db(self) -> None:
        """
        Check DD validity against the source database(s).

        Also caches SQLAlchemy source column types.
        """
        for d in self.get_source_databases():
            db = self.config.sources[d]

            for t in self.get_src_tables(d):

                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                # Ensure each source table maps to only one destination table
                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                dt = self.get_dest_tables_for_src_db_table(d, t)
                if len(dt) > 1:
                    raise ValueError(
                        f"Source table {d}.{t} maps to >1 destination "
                        f"table: {', '.join(dt)}")

                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                # Ensure source table is in database
                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                if t not in db.table_names:
                    log.debug(
                        f"Source database {d!r} has tables: {db.table_names}")
                    raise ValueError(
                        f"Table {t!r} missing from source database {d!r}")

                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                # Row checks: preamble
                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                rows = self.get_rows_for_src_table(d, t)
                # We may need to cross-reference rows, so all rows need to know
                # their type.
                for r in rows:
                    if r.src_field not in db.metadata.tables[t].columns:
                        raise ValueError(
                            f"Column {r.src_field!r} missing from table {t!r} "
                            f"in source database {d!r}")
                    sqla_coltype = (
                        db.metadata.tables[t].columns[r.src_field].type
                    )
                    r.set_src_sqla_coltype(sqla_coltype)  # CACHES TYPE HERE

                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                # If PID field is required, is it present?
                # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                needs_pidfield = any(r.being_scrubbed and not r.omit
                                     for r in rows)
                # Before 2021-12-07, we used to check r.master_pid, too.
                # However, if nothing is being scrubbed, then the lack of a
                # link via primary PID is a researcher inconvenience, not an
                # de-identification risk.
                if needs_pidfield and not self.get_pid_name(d, t):
                    raise ValueError(
                        f"Source table {d}.{t} has a "
                        f"{AlterMethodType.SCRUBIN.value!r} "
                        f"field but no primary patient ID field"
                    )

                n_pks = 0
                for r in rows:
                    # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                    # Data types for special rows
                    # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                    if r.primary_pid:
                        ensure_no_source_type_mismatch(r, self.config.pidtype,
                                                       primary_pid=True)
                    if r.master_pid:
                        ensure_no_source_type_mismatch(r, self.config.mpidtype,
                                                       primary_pid=False)

                    # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                    # Too many PKs?
                    # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                    if r.pk:
                        n_pks += 1
                        if n_pks > 1:
                            raise ValueError(
                                f"Table {d}.{t} has >1 source PK set")

                    # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                    # Duff alter method?
                    # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
                    for am in r.alter_methods:
                        if am.extract_from_blob:
                            extrow = next(
                                (r2 for r2 in rows
                                    if r2.src_field == am.extract_ext_field),
                                None
                            )
                            if extrow is None:
                                raise ValueError(
                                    f"alter_method = {r.alter_method}, "
                                    f"but field {am.extract_ext_field} "
                                    f"not found in the same table")
                            if not is_sqlatype_text_over_one_char(
                                    extrow.src_sqla_coltype):
                                raise ValueError(
                                    f"alter_method = {r.alter_method}, but "
                                    f"field {am.extract_ext_field}, which "
                                    f"should contain an extension or "
                                    f"filename, is not text of >1 character")

    def check_valid(self,
                    prohibited_fieldnames: List[str] = None,
                    check_against_source_db: bool = True) -> None:
        """
        Check DD validity, internally ± against the source database(s).

        Args:
            prohibited_fieldnames:
                list of prohibited destination fieldnames
            check_against_source_db:
                check validity against the source database(s)?

        Raises:
            :exc:`ValueError` if the DD is invalid
        """
        if prohibited_fieldnames is None:
            prohibited_fieldnames = []  # type: List[str]
        log.info("Checking data dictionary...")
        if not self.rows:
            raise ValueError("Empty data dictionary")
        if not self.get_dest_tables():
            raise ValueError("Empty data dictionary after removing "
                             "redundant tables")

        # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
        # Check (or re-check) individual rows
        # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
        log.debug("Checking DD: individual row validity...")
        for r in self.rows:
            r.check_valid()

        # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
        # Check collective consistency
        # ~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~

        log.debug("Checking DD: prohibited flags...")
        for d in self.get_source_databases():
            for t in self.get_src_tables(d):
                # This will have excluded all tables where all rows are
                # omitted. So now we have only active tables, for which we
                # cannot combine certain flags.
                # (We used to prohibit these combinations at all times, in the
                # DataDictionaryRow class, but it's inconvenient to have to
                # alter these flags if you want to omit the whole table.)
                for r in self.get_rows_for_src_table(d, t):
                    if r.add_src_hash and r.omit:
                        raise ValueError(
                            f"Do not set omit on "
                            f"src_flags={SrcFlag.ADD_SRC_HASH} fields -- "
                            f"currently set for {r.src_signature}")
                    if r.constant and r.omit:
                        raise ValueError(
                            f"Do not set omit on "
                            f"src_flags={SrcFlag.CONSTANT} fields -- "
                            f"currently set for {r.src_signature}")

        log.debug("Checking DD: table consistency...")
        for d, t in self.get_scrub_from_db_table_pairs():
            pid_field = self.get_pid_name(d, t)
            if not pid_field:
                raise ValueError(
                    f"Scrub-source table {d}.{t} must have a patient ID field "
                    f"(one with flag {SrcFlag.PRIMARY_PID})"
                )

        log.debug("Checking DD: prohibited fieldnames...")
        if prohibited_fieldnames:
            for r in self.rows:
                r.check_prohibited_fieldnames(prohibited_fieldnames)

        log.debug("Checking DD: opt-out fields...")
        for t in self.get_optout_defining_fields():
            (src_db, src_table, optout_colname, pid_colname, mpid_colname) = t
            if not pid_colname and not mpid_colname:
                raise ValueError(
                    f"Field {src_db}.{src_table}.{optout_colname} has "
                    f"src_flags={SrcFlag.OPT_OUT} set, but that table does "
                    f"not have a primary patient ID field or a master patient "
                    f"ID field")

        log.debug("Checking DD: destination tables...")
        for t in self.get_dest_tables():
            sdt = self.get_src_dbs_tables_for_dest_table(t)
            if len(sdt) > 1:
                raise ValueError(
                    "Destination table {t} is mapped to by multiple source "
                    "databases: {s}".format(
                        t=t,
                        s=", ".join(["{}.{}".format(s[0], s[1]) for s in sdt]),
                    ))

        log.debug("Checking DD: duplicate source rows?")
        src_sigs = [r.src_signature for r in self.rows]
        src_duplicates = [
            item for item, count in Counter(src_sigs).items()
            if count > 1
        ]
        if src_duplicates:
            raise ValueError(f"Duplicate source rows: {src_duplicates}")

        log.debug("Checking DD: duplicate destination rows?")
        dst_sigs = [r.dest_signature for r in self.rows if not r.omit]
        dst_duplicates = [
            item for item, count in Counter(dst_sigs).items()
            if count > 1
        ]
        if dst_duplicates:
            raise ValueError(f"Duplicate destination rows: {dst_duplicates}")

        if check_against_source_db:
            log.debug("Checking DD against source database tables...")
            self.check_against_source_db()

        log.debug("Checking DD: global patient-defining fields...")
        n_definers = sum([1 if x.defines_primary_pids else 0
                          for x in self.rows])
        if n_definers == 0:
            if self.config.allow_no_patient_info:
                log.warning("NO PATIENT-DEFINING FIELD! DATABASE(S) WILL "
                            "BE COPIED, NOT ANONYMISED.")
            else:
                raise ValueError(
                    f"No patient-defining field! (And "
                    f"{AnonymiseConfigKeys.ALLOW_NO_PATIENT_INFO} not set.)"
                )
        elif n_definers > 1:
            log.warning(
                f"Unusual: >1 field with "
                f"src_flags={SrcFlag.DEFINES_PRIMARY_PIDS} set.")

        log.debug("... DD checked.")

    # -------------------------------------------------------------------------
    # Saving
    # -------------------------------------------------------------------------

    def write(self, filename: str, filetype: str = None) -> None:
        """
        Writes the dictionary, either specifying the filetype or autodetecting
        it from the specified filename.

        Args:
            filename:
                Name of file to write, or "-" for stdout (in which case the
                filetype is forced to TSV).
            filetype:
                File type as one of ``.ods``, ``.tsv``, or ``.xlsx``;
                alternatively, use ``None`` to autodetect from the filename.
        """
        if filename == "-":
            with smart_open(filename, "wt") as f:  # type: TextIO
                self.write_tsv_file(f)
            return
        ext = filetype or os.path.splitext(filename)[1]
        if ext == EXT_TSV:
            self.write_tsv(filename)
        elif ext == EXT_ODS:
            self.write_ods(filename)
        elif ext == EXT_XLSX:
            self.write_xlsx(filename)
        else:
            raise ValueError(f"Unknown data dictionary extension: {ext!r}")

    def get_tsv(self) -> str:
        """
        Return the DD in TSV format.
        """
        return "\n".join(
            ["\t".join(DataDictionaryRow.ROWNAMES)] +
            [r.get_tsv() for r in self.rows]
        )

    def write_tsv_file(self, file: TextIO) -> None:
        """
        Writes the dictionary to a TSV file.
        """
        file.write(self.get_tsv())

    def write_tsv(self, filename: str) -> None:
        """
        Writes the dictionary to a TSV file.
        """
        log.info(f"Saving data dictionary as TSV: {filename}")
        with open(filename, "wt") as f:
            self.write_tsv_file(f)

    def _as_dict(self) -> Dict[str, Any]:
        """
        Returns an ordered dictionary representation used for writing
        spreadsheets.
        """
        sheetname = "data_dictionary"
        rows = [
            DataDictionaryRow.header_row()
        ] + [
            ddr.as_row() for ddr in self.rows
        ]
        data = OrderedDict()
        data[sheetname] = rows
        return data

    def write_ods(self, filename: str) -> None:
        """
        Writes the dictionary to an OpenOffice spreadsheet (ODS) file.
        """
        log.info(f"Saving data dictionary as ODS: {filename}")
        pyexcel_ods.save_data(filename, self._as_dict())

    def write_xlsx(self, filename: str) -> None:
        """
        Writes the dictionary to an Excel (XLSX) file.
        """
        log.info(f"Saving data dictionary as XLSX: {filename}")
        pyexcel_xlsx.save_data(filename, self._as_dict())

    # -------------------------------------------------------------------------
    # Global DD queries
    # -------------------------------------------------------------------------

    @lru_cache(maxsize=None)
    def get_source_databases(self) -> AbstractSet[str]:
        """
        Return a SortedSet of source database names.
        """
        return SortedSet([
             ddr.src_db
             for ddr in self.rows
             if ddr.required
         ])

    @lru_cache(maxsize=None)
    def get_scrub_from_db_table_pairs(self) -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of ``source_database_name, source_table`` tuples
        where those fields contain ``scrub_src`` (scrub-from) information.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table)
            for ddr in self.rows
            if ddr.scrub_src
        ])
        # even if omit flag set

    @lru_cache(maxsize=None)
    def get_src_db_tablepairs(self) -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of all ``source_database_name, source_table``
        tuples.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table)
            for ddr in self.rows
        ])

    @lru_cache(maxsize=None)
    def get_src_db_tablepairs_w_pt_info(self) -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of ``source_database_name, source_table`` tuples
        for tables that contain patient information.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table)
            for ddr in self.rows
            if ddr.contains_patient_info
        ])

    def get_src_db_tablepairs_w_no_pt_info(self) \
            -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of ``source_database_name, source_table`` tuples
        for tables that contain no patient information.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table)
            for ddr in self.rows
            if not ddr.contains_patient_info
        ])

    def get_tables_w_no_pt_info(self) -> AbstractSet[str]:
        """
        Return a SortedSet of ``source_table`` names for tables that contain no
        patient information.
        """
        tables_with_pt_info = SortedSet([
            ddr.src_table
            for ddr in self.rows
            if ddr.contains_patient_info
        ])
        all_tables = SortedSet([ddr.src_table for ddr in self.rows])
        return all_tables - tables_with_pt_info

    def get_tables_w_scrub_src(self) -> AbstractSet[str]:
        """
        Return a SortedSet of ``source_table`` names for tables that contain
        ``scrub_src`` information, i.e. that contribute to anonymisation.
        """
        return SortedSet([
            ddr.src_table
            for ddr in self.rows
            if ddr.contains_scrub_src
        ])

    @lru_cache(maxsize=None)
    def get_src_db_tablepairs_w_int_pk(self) -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of ``source_database_name, source_table`` tuples
        for tables that have an integer PK.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table)
            for ddr in self.rows
            if self.get_int_pk_ddr(ddr.src_db, ddr.src_table) is not None
        ])

    @lru_cache(maxsize=None)
    def get_src_dbs_tables_with_no_pt_info_no_pk(self) \
            -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of ``source_database_name, source_table`` tuples
        where the table has no patient information and no integer PK.
        """
        return (
            self.get_src_db_tablepairs() -
            self.get_src_db_tablepairs_w_pt_info() -
            self.get_src_db_tablepairs_w_int_pk()
        )

    @lru_cache(maxsize=None)
    def get_src_dbs_tables_with_no_pt_info_int_pk(self) \
            -> AbstractSet[Tuple[str, str]]:
        """
        Return a SortedSet of ``source_database_name, source_table`` tuples
        where the table has no patient information and has an integer PK.
        """
        return (
            (self.get_src_db_tablepairs() -
                self.get_src_db_tablepairs_w_pt_info()) &  # & is intersection
            self.get_src_db_tablepairs_w_int_pk()
        )

    @lru_cache(maxsize=None)
    def get_dest_tables(self) -> AbstractSet[str]:
        """
        Return a SortedSet of all destination table names.
        """
        return SortedSet([
            ddr.dest_table
            for ddr in self.rows
            if not ddr.omit
        ])

    @lru_cache(maxsize=None)
    def get_dest_tables_with_patient_info(self) -> AbstractSet[str]:
        """
        Return a SortedSet of destination table names that have patient
        information.
        """
        return SortedSet([
            ddr.dest_table
            for ddr in self.rows
            if ddr.contains_patient_info and not ddr.omit
        ])

    @lru_cache(maxsize=None)
    def get_optout_defining_fields(self) \
            -> AbstractSet[Tuple[str, str, str, str, str]]:
        """
        Return a SortedSet of ``src_db, src_table, src_field, pidfield,
        mpidfield`` tuples for rows that define opt-out information.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table, ddr.src_field,
                self.get_pid_name(ddr.src_db, ddr.src_table),
                self.get_mpid_name(ddr.src_db, ddr.src_table))
            for ddr in self.rows
            if ddr.opt_out_info
        ])

    @lru_cache(maxsize=None)
    def get_mandatory_scrubber_sigs(self) -> AbstractSet[str]:
        """
        Return a set of field signatures (strings of the format
        ``db.table.column``) for all rows representing "required scrubber"
        fields -- that is, rows that must have at least one non-NULL value for
        each patient, or the patient won't get processed.
        """
        return set([ddr.src_signature for ddr in self.rows
                    if ddr.required_scrubber])

    # -------------------------------------------------------------------------
    # Queries by source DB
    # -------------------------------------------------------------------------

    @lru_cache(maxsize=None)
    def get_src_tables(self, src_db: str) -> AbstractSet[str]:
        """
        For a given source database name, return a SortedSet of all source
        tables that are required (that is, ones being copied and ones providing
        vital patient information).
        """
        return SortedSet([
            ddr.src_table
            for ddr in self.rows
            if ddr.src_db == src_db and ddr.required
        ])

    @lru_cache(maxsize=None)
    def get_src_tables_with_active_dest(self, src_db: str) -> AbstractSet[str]:
        """
        For a given source database name, return a SortedSet of its source
        tables that have an active destination.
        """
        return SortedSet([
            ddr.src_table
            for ddr in self.rows
            if ddr.src_db == src_db and not ddr.omit
        ])

    @lru_cache(maxsize=None)
    def get_src_tables_with_patient_info(self, src_db: str) -> AbstractSet[str]:
        """
        For a given source database name, return a SortedSet of source tables
        that have patient information.
        """
        return SortedSet([
            ddr.src_table
            for ddr in self.rows
            if ddr.src_db == src_db and ddr.contains_patient_info
        ])

    @lru_cache(maxsize=None)
    def get_patient_src_tables_with_active_dest(self, src_db: str) \
            -> AbstractSet[str]:
        """
        For a given source database name, return a SortedSet of source tables
        that contain patient information and have an active destination table.
        """
        return (
            self.get_src_tables_with_active_dest(src_db) &
            self.get_src_tables_with_patient_info(src_db)
        )

    # -------------------------------------------------------------------------
    # Queries by source DB/table
    # -------------------------------------------------------------------------

    @lru_cache(maxsize=None)
    def get_dest_tables_for_src_db_table(
            self, src_db: str, src_table: str) -> AbstractSet[str]:
        """
        For a given source database/table, return a SortedSet of destination
        tables.
        """
        return SortedSet([
            ddr.dest_table
            for ddr in self.rows
            if (ddr.src_db == src_db and
                ddr.src_table == src_table and
                not ddr.omit)
        ])

    @lru_cache(maxsize=None)
    def get_dest_table_for_src_db_table(
            self, src_db: str, src_table: str) -> str:
        """
        For a given source database/table, return the single or the first
        destination table.
        """
        return list(
            self.get_dest_tables_for_src_db_table(src_db, src_table))[0]

    @lru_cache(maxsize=None)
    def get_rows_for_src_table(self, src_db: str, src_table: str) \
            -> AbstractSet[DataDictionaryRow]:
        """
        For a given source database name/table, return a SortedSet of DD rows.
        """
        return SortedSet([
            ddr
            for ddr in self.rows
            if ddr.src_db == src_db and ddr.src_table == src_table
        ])

    @lru_cache(maxsize=None)
    def get_fieldnames_for_src_table(self, src_db: str, src_table: str) \
            -> AbstractSet[DataDictionaryRow]:
        """
        For a given source database name/table, return a SortedSet of source
        fields.
        """
        return SortedSet([
            ddr.src_field
            for ddr in self.rows
            if ddr.src_db == src_db and ddr.src_table == src_table
        ])

    @lru_cache(maxsize=None)
    def get_scrub_from_rows(self, src_db: str, src_table: str) \
            -> AbstractSet[DataDictionaryRow]:
        """
        Return a SortedSet of DD rows for all fields containing ``scrub_src``
        (scrub-from) information.
        """
        return SortedSet([
            ddr
            for ddr in self.rows
            if (ddr.scrub_src and
                ddr.src_db == src_db and
                ddr.src_table == src_table)
        ])
        # even if omit flag set

    def get_scrub_from_rows_as_fieldinfo(
            self,
            src_db: str,
            src_table: str,
            depth: int,
            max_depth: int) -> List[ScrubSourceFieldInfo]:
        """
        Using :meth:`get_scrub_from_rows`, as a list of
        :class:`ScrubSourceFieldInfo` objects, which is more convenient for
        scrubbing.

        Args:
            src_db:
                Source database name.
            src_table:
                Source table name.
            depth:
                Current recursion depth for looking up third-party information.
            max_depth:
                Maximum permitted recursion depth for looking up third-party
                information.
        """
        ddrows = self.get_scrub_from_rows(src_db, src_table)
        infolist = []  # type: List[ScrubSourceFieldInfo]
        for ddr in ddrows:
            info = ScrubSourceFieldInfo(
                is_mpid=(
                    depth == 0 and ddr.master_pid
                    # The check for "depth == 0" means that third-party
                    # information is never marked as patient-related.
                ),
                is_patient=(
                    depth == 0 and ddr.scrub_src is ScrubSrc.PATIENT
                ),
                recurse=(
                    depth < max_depth
                    and ddr.scrub_src is ScrubSrc.THIRDPARTY_XREF_PID
                ),
                required_scrubber=ddr.required_scrubber,
                scrub_method=PersonalizedScrubber.get_scrub_method(
                    ddr.src_datatype,
                    ddr.scrub_method
                ),
                signature=ddr.src_signature,
                value_fieldname=ddr.src_field,
            )
            infolist.append(info)
        return infolist

    @lru_cache(maxsize=None)
    def get_pk_ddr(self, src_db: str, src_table: str) \
            -> Optional[DataDictionaryRow]:
        """
        For a given source database name and table, return the DD row for the
        PK for that table, whether integer or not.

        Will return ``None`` if no such data dictionary row exists.
        """
        for ddr in self.rows:
            if (ddr.src_db == src_db and
                    ddr.src_table == src_table and
                    ddr.pk):
                return ddr
        return None

    @lru_cache(maxsize=None)
    def get_int_pk_ddr(self, src_db: str, src_table: str) \
            -> Optional[DataDictionaryRow]:
        """
        For a given source database name and table, return the DD row for the
        integer PK for that table.

        Will return ``None`` if no such data dictionary row exists.
        """
        for ddr in self.rows:
            if (ddr.src_db == src_db and
                    ddr.src_table == src_table and
                    ddr.pk and
                    is_sqltype_integer(ddr.src_datatype)):
                return ddr
        return None

    @lru_cache(maxsize=None)
    def get_int_pk_name(self, src_db: str, src_table: str) -> Optional[str]:
        """
        For a given source database name and table, return the field name of
        the integer PK for that table (or ``None`` if there isn't one).
        """
        ddr = self.get_int_pk_ddr(src_db, src_table)
        if ddr is None:
            return None
        return ddr.src_field

    @lru_cache(maxsize=None)
    def has_active_destination(self, src_db: str, src_table: str) -> bool:
        """
        For a given source database name and table, does it have an active
        destination?
        """
        for ddr in self.rows:
            if (ddr.src_db == src_db and
                    ddr.src_table == src_table and
                    not ddr.omit):
                return True
        return False

    @lru_cache(maxsize=None)
    def get_pid_name(self, src_db: str, src_table: str) -> Optional[str]:
        """
        For a given source database name and table: return the field name of
        the field providing primary PID information (or ``None`` if there isn't
        one).
        """
        for ddr in self.rows:
            if (ddr.src_db == src_db and
                    ddr.src_table == src_table and
                    ddr.primary_pid):
                return ddr.src_field
        return None

    @lru_cache(maxsize=None)
    def get_mpid_name(self, src_db: str, src_table: str) -> Optional[str]:
        """
        For a given source database name and table: return the field name of
        the field providing master PID (MPID) information (or ``None`` if there
        isn't one).
        """
        for ddr in self.rows:
            if (ddr.src_db == src_db and
                    ddr.src_table == src_table and
                    ddr.master_pid):
                return ddr.src_field
        return None

    # -------------------------------------------------------------------------
    # Queries by destination table
    # -------------------------------------------------------------------------

    @lru_cache(maxsize=None)
    def get_src_dbs_tables_for_dest_table(
            self, dest_table: str) -> AbstractSet[Tuple[str, str]]:
        """
        For a given destination table, return a SortedSet of ``dbname, table``
        tuples.
        """
        return SortedSet([
            (ddr.src_db, ddr.src_table)
            for ddr in self.rows
            if ddr.dest_table == dest_table
        ])

    @lru_cache(maxsize=None)
    def get_rows_for_dest_table(
            self, dest_table: str) -> AbstractSet[DataDictionaryRow]:
        """
        For a given destination table, return a SortedSet of DD rows.
        """
        return SortedSet([
            ddr
            for ddr in self.rows
            if ddr.dest_table == dest_table and not ddr.omit
        ])

    # -------------------------------------------------------------------------
    # SQLAlchemy Table objects
    # -------------------------------------------------------------------------

    @lru_cache(maxsize=None)
    def get_dest_sqla_table(self, tablename: str,
                            timefield: str = None,
                            add_mrid_wherever_rid_added: bool = False) -> Table:
        """
        For a given destination table name, return an
        :class:`sqlalchemy.sql.schema.Table` object for the destination table
        (which we will create).
        """
        metadata = self.config.destdb.metadata
        columns = []  # type: List[Column]
        for ddr in self.get_rows_for_dest_table(tablename):
            columns.append(ddr.dest_sqla_column)
            if ddr.add_src_hash:
                columns.append(self._get_srchash_sqla_column())
            if ddr.primary_pid:
                columns.append(self._get_trid_sqla_column())
                if add_mrid_wherever_rid_added:
                    columns.append(self._get_mrid_sqla_column())
        if timefield:
            timecol = Column(timefield, DateTime)
            columns.append(timecol)
        return Table(tablename, metadata, *columns, **TABLE_KWARGS)

    def _get_srchash_sqla_column(self) -> Column:
        """
        Returns a :class:`sqlalchemy.sql.schema.Column` object for the
        "source hash" column (which is inserted into many destination tables
        so they can record the hash of their source, for change detection).
        """
        return Column(
            self.config.source_hash_fieldname,
            self.config.sqltype_encrypted_pid,
            comment='Hashed amalgamation of all source fields'
        )

    def _get_trid_sqla_column(self) -> Column:
        """
        Returns a :class:`sqlalchemy.sql.schema.Column` object for the "TRID"
        column. This is inserted into all patient-related destination tables as
        a high-speed (integer) but impermanent research ID -- a transient
        research ID (TRID).
        """
        return Column(
            self.config.trid_fieldname,
            TridType,
            nullable=False,
            comment='Transient integer research ID (TRID)'
        )

    def _get_mrid_sqla_column(self) -> Column:
        """
        Returns a :class:`sqlalchemy.sql.schema.Column` object for the "MRID"
        column. This is inserted into all patient-related destination tables
        where the flag 'add_mrid_wherever_rid_added' is set.
        """
        return Column(
            self.config.master_research_id_fieldname,
            self.config.sqltype_encrypted_pid,
            nullable=True,
            comment='Master research ID (MRID)'
        )

    # -------------------------------------------------------------------------
    # Clear caches
    # -------------------------------------------------------------------------

    def cached_funcs(self) -> List[Any]:
        """
        Returns a list of our methods that are cached. See
        :func:`clear_caches`.
        """
        return [
            self.get_source_databases,
            self.get_scrub_from_db_table_pairs,
            self.get_src_db_tablepairs,
            self.get_src_db_tablepairs_w_pt_info,
            self.get_src_db_tablepairs_w_int_pk,
            self.get_src_dbs_tables_with_no_pt_info_no_pk,
            self.get_src_dbs_tables_with_no_pt_info_int_pk,
            self.get_dest_tables,
            self.get_dest_tables_with_patient_info,
            self.get_optout_defining_fields,
            self.get_mandatory_scrubber_sigs,

            self.get_src_tables,
            self.get_src_tables_with_active_dest,
            self.get_src_tables_with_patient_info,
            self.get_patient_src_tables_with_active_dest,

            self.get_dest_tables_for_src_db_table,
            self.get_dest_table_for_src_db_table,
            self.get_rows_for_src_table,
            self.get_fieldnames_for_src_table,
            self.get_scrub_from_rows,
            self.get_pk_ddr,
            self.get_int_pk_ddr,
            self.get_int_pk_name,
            self.has_active_destination,
            self.get_pid_name,
            self.get_mpid_name,

            self.get_src_dbs_tables_for_dest_table,
            self.get_rows_for_dest_table,

            self.get_dest_sqla_table,
        ]

    def clear_caches(self) -> None:
        """
        Clear all our cached information.
        """
        for func in self.cached_funcs():
            func.cache_clear()

    def debug_cache_hits(self) -> None:
        """
        Report cache hit information for our caches, to the Python log.
        """
        for func in self.cached_funcs():
            log.debug(f"{func.__name__}: {func.cache_info()}")

    # -------------------------------------------------------------------------
    # Filtering
    # -------------------------------------------------------------------------

    KEEP_FUNCTION_TYPE = Callable[[DataDictionaryRow], bool]
    # ... returns keep (True/False)

    def remove_rows_by_filter(self, keep: KEEP_FUNCTION_TYPE) -> None:
        """
        Removes any rows that do not pass a filter function.

        Args:
            keep:
                Function taking a data dictionary row as an argument, and
                returning a boolean of whether to keep the row.
        """
        self.rows = [
            row for row in self.rows if keep(row)
        ]

    def omit_rows_by_filter(self, keep: KEEP_FUNCTION_TYPE) -> None:
        """
        Set to "omit" any rows that do not pass a filter function.
        Does not alter any rows already set to omit.

        Args:
            keep:
                Function taking a data dictionary row as an argument, and
                returning a boolean of whether to keep the row.
        """
        for row in self.rows:
            if not row.omit:
                row.omit = not keep(row)

    MODIFYING_KEEP_FUNCTION_TYPE = Callable[[DataDictionaryRow],
                                            Optional[DataDictionaryRow]]
    # returns the row (perhaps modified) to keep, or None to reject

    def remove_rows_by_modifying_filter(
            self, keep_modify: MODIFYING_KEEP_FUNCTION_TYPE) -> None:
        """
        Removes any rows that do not pass a filter function; allows the filter
        function to modify rows that are kept.

        Args:
            keep_modify:
                Function taking a data dictionary row as an argument, and
                returning either the row (potentially modified) to retain it,
                or ``None`` to reject it.
        """
        new_rows = []  # type: List[DataDictionaryRow]
        for row in self.rows:
            result = keep_modify(row)
            if result is not None:
                new_rows.append(result)
        self.rows = new_rows
